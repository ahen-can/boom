"""End-to-end tests for the deck builder. Run: python -m pytest deck-builder/tests -q"""
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "skill" / "pwc-deck-builder"
SCRIPTS = SKILL / "scripts"
sys.path.insert(0, str(SCRIPTS))

import build_deck  # noqa: E402
import check_deck  # noqa: E402
import deckkit as dk  # noqa: E402

PATTERNS = SKILL / "assets" / "patterns.pptx"
ICONS = SKILL / "assets" / "icons"
PLAN = json.loads((ROOT / "tests" / "sample_plan.json").read_text())


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("build")
    result = build_deck.build(PLAN, PATTERNS, ICONS, out)
    assert result["ok"], result
    return result, Presentation(result["output"])


def all_text(slide):
    return "\n".join(dk.shape_text(s) for s, _ in dk.iter_shapes(slide.shapes))


def test_slide_count_and_order(built):
    _, prs = built
    assert len(prs.slides) == len(PLAN["slides"])
    titles = [all_text(s) for s in prs.slides]
    for spec, text in zip(PLAN["slides"], titles):
        assert spec["fill"].get("title", spec["fill"].get("section_title")) in text


def test_no_leftover_tokens(built):
    _, prs = built
    for s in prs.slides:
        assert "{{" not in all_text(s)


def test_unused_item_removed(built):
    _, prs = built
    slide = prs.slides[4]  # OBJ_NUMBERED_4 with three items
    names = {s.name for s, _ in dk.iter_shapes(slide.shapes)}
    assert "item3" in names and "item4" not in names
    assert "item4_head" not in names


def test_list_values_become_paragraphs(built):
    _, prs = built
    body = dk.find_shape(prs.slides[5], "item1_body")
    assert [p.text for p in body.text_frame.paragraphs] == PLAN["slides"][5]["fill"]["item1_body"]


def test_icons_swapped(built):
    result, prs = built
    assert any("item2_icon ← data.png" in line for line in result["log"])
    slide = prs.slides[3]
    pic = dk.find_shape(slide, "item2_icon")
    assert pic is not None and pic.image.size == (256, 256)


def test_relationships_resolve(built):
    """Every r:id/r:embed in every slide points at an existing relationship."""
    result, _ = built
    with zipfile.ZipFile(result["output"]) as z:
        names = z.namelist()
        assert len(names) == len(set(names))
        for n in names:
            m = re.fullmatch(r"ppt/slides/(slide\d+)\.xml", n)
            if not m:
                continue
            xml = z.read(n).decode()
            rels = z.read(f"ppt/slides/_rels/{m.group(1)}.xml.rels").decode()
            rids = set(re.findall(r'Id="(rId\d+)"', rels))
            used = set(re.findall(r'r:(?:id|embed|link)="(rId\d+)"', xml))
            assert used <= rids, (n, used - rids)


def test_check_is_clean(built):
    result, _ = built
    assert result["check"]["errors"] == 0, result["check"]["issues"]


def test_check_catches_problems(tmp_path):
    plan = json.loads(json.dumps(PLAN))
    plan["slides"][3]["fill"]["item1_body"] = "We leverage synergies. " + "Long text that keeps going. " * 40
    plan["slides"].insert(4, json.loads(json.dumps(plan["slides"][3])))
    result = build_deck.build(plan, PATTERNS, ICONS, tmp_path)
    issues = [i["issue"] for i in result["check"]["issues"]]
    assert any("banned phrase: 'leverage'" in i for i in issues)
    assert any("overflows" in i for i in issues)
    assert any("same pattern as previous slide" in i for i in issues)


def test_validation_errors(tmp_path):
    plan = {"slides": [{"pattern": "NOPE", "fill": {}}, {"pattern": "OBJ_3COL", "fill": {"bogus": "x"}}]}
    result = build_deck.build(plan, PATTERNS, ICONS, tmp_path)
    assert not result["ok"]
    assert any("unknown pattern" in e for e in result["errors"])
    assert any("unknown slots ['bogus']" in e for e in result["errors"])


def test_split_token_is_filled(tmp_path):
    prs = Presentation(str(PATTERNS))
    slide = prs.slides[5]  # TWO_COL
    shape = dk.find_shape(slide, "left_head")
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = "{{left_"
    r = p.add_run()
    r.text = "head}}"
    filled, unfilled = dk.fill_tokens(slide, {"left_head": "Joined"})
    assert shape.text_frame.text == "Joined" and "left_head" in filled


def test_prep_patterns_round_trip(tmp_path, built):
    result, _ = built
    deck = tmp_path / "human.pptx"
    shutil.copy(result["output"], deck)
    subprocess.run([sys.executable, str(SCRIPTS / "prep_patterns.py"), str(deck), "--out", str(tmp_path)],
                   check=True, capture_output=True)
    prs = Presentation(str(tmp_path / "human_tokenised.pptx"))
    toks = dk.slide_tokens(prs.slides[3])
    assert {"title", "tracker", "source", "item1_head", "item1_body", "item3_body"} <= set(toks)
    names = {s.name for s, _ in dk.iter_shapes(prs.slides[3].shapes)}
    assert {"item1_icon", "item2_icon", "item3_icon"} <= names
    assert dk.pattern_id(prs.slides[3]) == "SLIDE_04"


def test_storylines(tmp_path, built):
    result, _ = built
    out = tmp_path / "s.md"
    subprocess.run([sys.executable, str(SCRIPTS / "extract_storylines.py"), result["output"], "--out", str(out)],
                   check=True, capture_output=True)
    text = out.read_text()
    assert "Three outcomes will tell Northwind the programme worked" in text
    assert "[Section] What we heard" in text


def test_catalogue_lists_patterns():
    text = subprocess.run([sys.executable, str(SCRIPTS / "inventory.py"), str(PATTERNS), "--catalogue"],
                          check=True, capture_output=True, text=True).stdout
    for pid in ("TITLE", "OBJ_3COL", "APPROACH_CHEVRON_4"):
        assert f"## {pid}" in text


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="no LibreOffice")
def test_render(tmp_path, built):
    result, _ = built
    out = subprocess.run([sys.executable, str(SCRIPTS / "render.py"), result["output"], "--out", str(tmp_path)],
                         check=True, capture_output=True, text=True, timeout=300).stdout
    data = json.loads(out)
    assert len(data["slides"]) == len(PLAN["slides"]) and Path(data["contact_sheet"]).is_file()
