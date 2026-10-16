"""Turn human-made slides into pattern slides: propose {{tokens}}, name shapes, add PATTERN notes.

    python prep_patterns.py my_slides.pptx [--out DIR] [--dry-run]

Writes <name>_tokenised.pptx and <name>_prep_report.md. Nothing is changed in the input file.
The guesses are positional (title at the top, source at the bottom, bold/larger text = item head,
text below or beside a head = its body), so always review the result in PowerPoint before using it.
"""
import argparse
import re
import statistics
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import deckkit as dk  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

EMU_PER_IN = 914400
SKIP_PH_TYPES = {"FOOTER", "SLIDE_NUMBER", "DATE"}
LABEL_RE = re.compile(r"^\s*([0-9]{1,2}|[A-Z]|[ivx]{1,4}|0[1-9])[.)]?\s*$", re.I)


def ph_type(shp):
    if getattr(shp, "is_placeholder", False) and shp.is_placeholder:
        return str(shp.placeholder_format.type).split(".")[-1].split(" ")[0]
    return None


def abs_box(shp, parents):
    """Shape box in slide coordinates (inches), accounting for group transforms."""
    x, y, w, h = shp.left or 0, shp.top or 0, shp.width or 0, shp.height or 0
    for g in reversed(parents):
        xfrm = g._element.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
        if xfrm is None:
            continue
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        choff, chext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
        if None in (off, ext, choff, chext):
            continue
        sx = int(ext.get("cx")) / max(1, int(chext.get("cx")))
        sy = int(ext.get("cy")) / max(1, int(chext.get("cy")))
        x = int(off.get("x")) + (x - int(choff.get("x"))) * sx
        y = int(off.get("y")) + (y - int(choff.get("y"))) * sy
        w, h = w * sx, h * sy
    return x / EMU_PER_IN, y / EMU_PER_IN, w / EMU_PER_IN, h / EMU_PER_IN


class Cand:
    def __init__(self, shp, parents):
        self.shp, self.parents = shp, parents
        self.x, self.y, self.w, self.h = abs_box(shp, parents)
        self.text = dk.shape_text(shp).strip()
        p = shp.text_frame.paragraphs[0] if shp.has_text_frame and shp.text_frame.paragraphs else None
        self.size = dk.effective_font_pt(shp, p.runs[0] if p is not None and p.runs else None, p)
        self.bold = bool(p is not None and p.runs and p.runs[0].font.bold)
        self.top_group = parents[0] if parents else None
        self.token = None

    @property
    def cx(self):
        return self.x + self.w / 2

    @property
    def cy(self):
        return self.y + self.h / 2


def reading_order(items, row_tol=0.4):
    return sorted(items, key=lambda c: (round(c.y / row_tol), c.x))


def assign_tokens(slide, slide_h):
    cands, notes = [], []
    for shp, parents in dk.iter_shapes(slide.shapes):
        if not (getattr(shp, "has_text_frame", False) and shp.has_text_frame):
            if getattr(shp, "has_table", False) and shp.has_table:
                notes.append(f"table '{shp.name}' left as is: replace its text by hand or with tokens")
            continue
        if ph_type(shp) in SKIP_PH_TYPES or not dk.shape_text(shp).strip() or dk.tokens_in(shp):
            continue
        c = Cand(shp, parents)
        if LABEL_RE.match(c.text):
            continue  # numbers/letters in badges stay as they are
        cands.append(c)
    if not cands:
        return [], notes

    # title
    titles = [c for c in cands if ph_type(c.shp) in ("TITLE", "CENTER_TITLE")]
    if not titles:
        top = [c for c in cands if c.y < slide_h * 0.25]
        titles = sorted(top, key=lambda c: (-c.size, c.y))[:1]
        if not titles:  # cover-style slides: the largest text anywhere, if it is clearly a heading
            biggest = max(cands, key=lambda c: (c.size, -c.y))
            titles = [biggest] if biggest.size >= 20 else []
    if titles:
        titles[0].token = "title"
    rest = [c for c in cands if c.token is None]

    # tracker (small text above the title) and source (small text at the bottom)
    title_y = titles[0].y if titles else slide_h * 0.1
    for c in rest:
        if c.y + c.h <= title_y + 0.05 and c.size <= 12:
            c.token = "tracker"
            break
    for c in reading_order([c for c in rest if c.token is None]):
        if (c.y > slide_h * 0.85 and c.size <= 11) or re.match(r"^(source|note)s?\b", c.text, re.I):
            c.token = "source" if c.token is None and "source" not in {x.token for x in cands} else c.token
    body = [c for c in cands if c.token is None]
    subtitles = [c for c in body if ph_type(c.shp) == "SUBTITLE"]
    for c in subtitles:
        c.token = "subtitle"
    body = [c for c in body if c.token is None]
    if not body:
        return cands, notes

    # items: one per top-level group when the body text sits in 2+ groups
    groups = {}
    for c in body:
        if c.top_group is not None:
            groups.setdefault(id(c.top_group), []).append(c)
    if len(groups) >= 2:
        ordered = reading_order([min(g, key=lambda c: (c.y, c.x)) for g in groups.values()])
        for n, first in enumerate(ordered, 1):
            members = sorted(groups[id(first.top_group)], key=lambda c: (c.y, c.x))
            first.top_group.name = f"item{n}"
            _head_body(members, f"item{n}")
        loose = [c for c in body if c.top_group is None]
        _loose(loose)
        return cands, notes

    # no groups: heads are bold or larger than the typical body text
    median = statistics.median(c.size for c in body)
    heads = [c for c in body if c.size > median + 0.5 or (c.bold and len(c.text) < 80)]
    if len(heads) >= 2:
        heads = reading_order(heads)
        for n, hd in enumerate(heads, 1):
            hd.token = f"item{n}_head"
        for c in body:
            if c.token:
                continue
            best = _nearest_head(c, heads)
            if best is not None:
                n = heads.index(best) + 1
                taken = sum(1 for x in body if x.token and x.token.startswith(f"item{n}_body"))
                c.token = f"item{n}_body" + (str(taken + 1) if taken else "")
    elif len(body) >= 2 and _parallel(body):
        for n, c in enumerate(reading_order(body), 1):
            c.token = f"item{n}_body"
    _loose([c for c in body if c.token is None])
    return cands, notes


def _head_body(members, prefix):
    if len(members) == 1:
        members[0].token = f"{prefix}_body"
        return
    head = max(members, key=lambda c: (c.size + (2 if c.bold else 0), -c.y))
    head.token = f"{prefix}_head"
    extra = 0
    for c in members:
        if c.token is None:
            extra += 1
            c.token = f"{prefix}_body" + (str(extra) if extra > 1 else "")


def _nearest_head(c, heads):
    """Head directly above (overlapping columns) or directly left on the same row."""
    best, best_d = None, 99
    for h in heads:
        overlap_x = min(c.x + c.w, h.x + h.w) - max(c.x, h.x)
        same_row = abs(c.cy - h.cy) < max(0.4, h.h)
        if overlap_x > 0.2 and c.y >= h.y:
            d = c.y - (h.y + h.h)
        elif same_row and c.x >= h.x + h.w - 0.1:
            d = c.x - (h.x + h.w)
        else:
            continue
        if -0.2 <= d < best_d and d < 2.0:
            best, best_d = h, d
    return best


def _parallel(body):
    """Similar-size boxes side by side or stacked: treat as items."""
    ws = [c.w for c in body]
    return max(ws) - min(ws) < 0.5 and len({round(c.y, 0) for c in body}) in (1, len(body))


def _loose(cands):
    for n, c in enumerate(reading_order(cands), 1):
        c.token = "body" if len(cands) == 1 else f"text{n}"


def name_item_shapes(slide, cands):
    """Rename pictures and background boxes that belong to an item: itemN_icon, itemN_bg."""
    heads = [c for c in cands if c.token and re.match(r"item\d+_(head|body)$", c.token)]
    renamed = []
    by_item = {}
    for c in heads:
        by_item.setdefault(c.token.split("_")[0], []).append(c)
    for shp, parents in dk.iter_shapes(slide.shapes):
        if shp.__class__.__name__ == "GroupShape" or dk.shape_text(shp).strip():
            continue
        x, y, w, h = abs_box(shp, parents)
        inside = [it for it, cs in by_item.items()
                  if any(x - 0.05 <= c.cx <= x + w + 0.05 and y - 0.05 <= c.cy <= y + h + 0.05 for c in cs)]
        if shp._element.tag == qn("p:pic"):
            near = [it for it, cs in by_item.items()
                    if any(abs((x + w / 2) - c.cx) < max(c.w / 2, 1.0) and 0 <= c.y - (y + h) < 1.0 for c in cs)
                    or any(abs((y + h / 2) - c.cy) < 0.6 and 0 <= c.x - (x + w) < 0.8 for c in cs)]
            if len(near) == 1:
                shp.name = f"{near[0]}_icon"
                renamed.append(shp.name)
        elif len(inside) == 1:
            shp.name = f"{inside[0]}_bg"
            renamed.append(shp.name)
    return renamed


def tokenise(c):
    tf = c.shp.text_frame
    paras = tf.paragraphs
    first = paras[0]
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    runs = first.runs
    if runs:
        runs[0].text = f"{{{{{c.token}}}}}"
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        first.text = f"{{{{{c.token}}}}}"
    if not dk.ITEM_RE.match(c.shp.name or "") and not c.shp.is_placeholder:
        c.shp.name = c.token


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--out", default=None)
    ap.add_argument("--dry-run", action="store_true", help="only write the report")
    a = ap.parse_args()
    src = Path(a.deck)
    out_dir = Path(a.out) if a.out else src.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = dk.open_deck(src)
    slide_h = prs.slide_height / EMU_PER_IN
    report = [f"# Pattern prep report: {src.name}", "",
              "Check each proposal in PowerPoint (Alt+F10 shows shape names). Fix tokens by editing the text,",
              "then fill in PURPOSE / USE WHEN / ITEMS in each slide's notes and give PATTERN a real name.", ""]
    for n, slide in enumerate(prs.slides, 1):
        cands, notes = assign_tokens(slide, slide_h)
        existing = dk.pattern_id(slide)
        report.append(f"## Slide {n}" + (f" ({existing})" if existing else ""))
        if not cands:
            report.append("- no text to tokenise\n")
            continue
        report.append("| token | shape | at (x, y) in | font | original text |")
        report.append("|---|---|---|---|---|")
        for c in reading_order(cands):
            snippet = c.text.replace("\n", " / ").replace("|", "/")[:70]
            report.append(f"| `{{{{{c.token}}}}}` | {c.shp.name} | {c.x:.1f}, {c.y:.1f} | {c.size:.0f}pt"
                          f"{' bold' if c.bold else ''} | {snippet} |")
        renamed = name_item_shapes(slide, cands)
        if renamed:
            report.append(f"\nRenamed shapes: {', '.join(renamed)}")
        for c in cands:
            tokenise(c)
        items = sorted({m.group(1) for c in cands if c.token and (m := dk.ITEM_RE.match(c.token))})
        if not existing:
            dk.set_notes(slide, f"PATTERN: SLIDE_{n:02d}\nPURPOSE: \nUSE WHEN: \n"
                                f"ITEMS: {('1-' + str(len(items))) if items else ''}")
            notes.append("speaker notes replaced with a PATTERN template (old notes removed)")
        pics = [shp.name for shp, _ in dk.iter_shapes(slide.shapes)
                if shp._element.tag == qn("p:pic") and "_icon" not in (shp.name or "")]
        if pics:
            notes.append(f"pictures kept as they are: {', '.join(pics)} (remove client logos/photos)")
        for note in notes:
            report.append(f"- {note}")
        report.append("")
    rep_path = out_dir / f"{src.stem}_prep_report.md"
    rep_path.write_text("\n".join(report), encoding="utf-8")
    print(f"wrote {rep_path}")
    if not a.dry_run:
        out = out_dir / f"{src.stem}_tokenised.pptx"
        prs.save(str(out))
        print(f"wrote {out}")


if __name__ == "__main__":
    main()
