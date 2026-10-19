"""Render a deck to PNGs (one per slide) and contact sheets, for visual checks.

    python render.py deck.pptx [--slides 1-15] [--out DIR] [--width 1280] [--timeout 150]
    python render.py deck.pdf [--slides 1-15] [--labels-from deck.pptx]
    python render.py --diagnose

Needs LibreOffice (soffice). PNGs come from PyMuPDF (pip-installed if missing), else pdftoppm.
Big decks: render in chunks (--slides 1-15, then 16-30, ...) so each run stays inside the
sandbox's time limit. Large pictures are shrunk in the temporary copy before conversion (the deck
itself is not changed). If LibreOffice still times out, export the deck to PDF in PowerPoint and
render the PDF instead: no LibreOffice needed, and --labels-from adds the pattern IDs. Contact sheets hold 12 slides each and are labelled with the slide number
and, for pattern libraries, the PATTERN ID.

Always prints one JSON object. On failure it has "ok": false, the stage that failed, the error
and a hint, and exits with code 1.
"""
import argparse
import importlib
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from PIL import Image, ImageDraw

sys.path.insert(0, str(Path(__file__).resolve().parent))
import deckkit as dk  # noqa: E402

PROFILE = Path(tempfile.gettempdir()) / "deckbuilder_lo_profile"  # reused: first start is the slow one


class RenderError(Exception):
    def __init__(self, stage, error, hint):
        super().__init__(error)
        self.stage, self.error, self.hint = stage, error, hint


def soffice_path():
    return shutil.which("soffice") or shutil.which("libreoffice")


def run_soffice(args, timeout):
    """Run LibreOffice in its own process group so a timeout kills every child it started."""
    soffice = soffice_path()
    if not soffice:
        raise RenderError("libreoffice", "soffice not found", "Check the deck in PowerPoint instead.")
    env = dict(os.environ, HOME=os.environ.get("HOME") or tempfile.gettempdir())
    cmd = [soffice, f"-env:UserInstallation=file://{PROFILE}", "--headless", "--norestore", *args]
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env, start_new_session=True)
    try:
        out, err = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        os.killpg(proc.pid, signal.SIGKILL)
        proc.communicate()
        raise RenderError("libreoffice", f"LibreOffice took longer than {timeout}s",
                          "Render fewer slides per run (--slides 1-5), or ask the user to export the deck to PDF in "
                          "PowerPoint and render the PDF with --labels-from <deck>.pptx.")
    return proc.returncode, out.decode(errors="ignore"), err.decode(errors="ignore")


def parse_slides(spec: str, total: int) -> list[int]:
    picked = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        nums = range(int(m.group(1)), int(m.group(2)) + 1) if m else [int(part)]
        picked += [n for n in nums if 1 <= n <= total and n not in picked]
    if not picked:
        raise RenderError("arguments", f"no slides in '{spec}' (deck has {total})", "Use e.g. --slides 1-15")
    return picked


def shrink_images(prs, max_px=1600, min_bytes=300_000) -> int:
    """Downscale big raster images in place (only in the render copy). Returns how many were shrunk."""
    import io

    done = 0
    for part in prs.part.package.iter_parts():
        ct = getattr(part, "content_type", "")
        if not ct.startswith("image/") or ct in ("image/svg+xml", "image/x-emf", "image/x-wmf"):
            continue
        blob = part.blob
        if len(blob) < min_bytes:
            continue
        try:
            im = Image.open(io.BytesIO(blob))
            fmt = im.format or "PNG"
            if max(im.size) <= max_px and len(blob) < 2 * min_bytes:
                continue
            im.thumbnail((max_px, max_px))
            buf = io.BytesIO()
            if fmt.upper() in ("JPEG", "JPG"):
                im.convert("RGB").save(buf, "JPEG", quality=80)
            else:
                im.save(buf, fmt)
            part._blob = buf.getvalue()
            done += 1
        except Exception:  # noqa: BLE001  (unreadable image: leave it)
            continue
    return done


def pattern_labels(pptx: Path | None, numbers: list[int]) -> list[str]:
    ids = {}
    if pptx:
        for n, s in enumerate(dk.open_deck(pptx).slides, 1):
            ids[n] = dk.pattern_id(s)
    return [f"Slide {n}" + (f"  {ids[n]}" if ids.get(n) else "") for n in numbers]


def subset_deck(deck: Path, slides: list[int], work: Path) -> tuple[Path, list[str]]:
    """Copy of the deck with only the chosen slides (unused media drops out). Returns path and labels."""
    prs = dk.open_deck(deck)
    all_slides = list(prs.slides)
    labels = []
    for n in slides:
        pid = dk.pattern_id(all_slides[n - 1])
        labels.append(f"Slide {n}" + (f"  {pid}" if pid else ""))
    keep = set(slides)
    for n, s in reversed(list(enumerate(all_slides, 1))):
        if n not in keep:
            dk.delete_slide(prs, s)
    shrink_images(prs)
    out = work / f"{deck.stem}_render_input.pptx"
    prs.save(str(out))
    return out, labels


def to_pdf(deck: Path, out: Path, timeout: int) -> Path:
    code, so, se = run_soffice(["--convert-to", "pdf", "--outdir", str(out), str(deck)], timeout)
    pdf = out / (deck.stem + ".pdf")
    if not pdf.is_file():
        raise RenderError("libreoffice", f"no PDF produced (exit {code}): {(se or so).strip()[-400:]}",
                          "Run render.py --diagnose and share its output.")
    return pdf


def load_pymupdf():
    try:
        return importlib.import_module("pymupdf")
    except ImportError:
        r = subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "pymupdf"],
                           capture_output=True, text=True, timeout=180)
        if r.returncode != 0:
            return None
        importlib.invalidate_caches()
        try:
            return importlib.import_module("pymupdf")
        except ImportError:
            return None


def pdf_to_pngs(pdf: Path, out: Path, width: int, numbers: list[int], pages: list[int] | None = None) -> list[Path]:
    """Render PDF pages to PNGs named after their slide numbers. `pages` (1-based) picks pages; default all."""
    fitz = load_pymupdf()
    if fitz is not None:
        paths = []
        with fitz.open(str(pdf)) as doc:
            chosen = [doc[p - 1] for p in pages] if pages else list(doc)
            for page, n in zip(chosen, numbers):
                zoom = width / page.rect.width
                p = out / f"slide{n:03d}.png"
                page.get_pixmap(matrix=fitz.Matrix(zoom, zoom)).save(str(p))
                paths.append(p)
        return paths
    if shutil.which("pdftoppm") and not pages:
        tmp = out / "_pages"
        tmp.mkdir(exist_ok=True)
        subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(width), "-scale-to-y", "-1", str(pdf),
                        str(tmp / "p")], check=True, timeout=240)
        paths = []
        for src, n in zip(sorted(tmp.glob("p*.png")), numbers):
            dest = out / f"slide{n:03d}.png"
            src.replace(dest)
            paths.append(dest)
        shutil.rmtree(tmp, ignore_errors=True)
        return paths
    raise RenderError("images", "neither PyMuPDF nor pdftoppm is available (pip install failed?)",
                      "The PDF was made; open it to check the slides.")


def contact_sheets(pngs, labels, out: Path, stem: str, per_sheet=12, cols=3, thumb_w=480) -> list[Path]:
    sheets = []
    for start in range(0, len(pngs), per_sheet):
        chunk = list(zip(pngs[start:start + per_sheet], labels[start:start + per_sheet]))
        thumbs = []
        for p, label in chunk:
            im = Image.open(p).convert("RGB")
            im.thumbnail((thumb_w, thumb_w))
            thumbs.append((im, label))
        th = max(t.height for t, _ in thumbs)
        rows = (len(thumbs) + cols - 1) // cols
        pad, lab = 16, 24
        sheet = Image.new("RGB", (cols * (thumb_w + pad) + pad, rows * (th + lab + pad) + pad), "white")
        d = ImageDraw.Draw(sheet)
        for i, (t, label) in enumerate(thumbs):
            x = pad + (i % cols) * (thumb_w + pad)
            y = pad + (i // cols) * (th + lab + pad)
            d.text((x, y), label, fill=(40, 40, 40))
            sheet.paste(t, (x, y + lab))
            d.rectangle([x - 1, y + lab - 1, x + t.width, y + lab + t.height], outline=(200, 200, 200))
        first = re.search(r"\d+", chunk[0][1]).group(0)
        last = re.search(r"\d+", chunk[-1][1]).group(0)
        dest = out / f"{stem}_contact_{int(first):03d}-{int(last):03d}.png"
        sheet.save(dest)
        sheets.append(dest)
    return sheets


def diagnose(timeout: int) -> dict:
    from pptx import Presentation

    info = {"soffice": soffice_path(), "profile": str(PROFILE)}
    if not info["soffice"]:
        info["result"] = "LibreOffice not installed"
        return info
    t = time.time()
    try:
        code, so, se = run_soffice(["--version"], timeout)
        info["version"] = (so or se).strip()[:120]
    except RenderError as e:
        info["version_error"] = e.error
    info["version_seconds"] = round(time.time() - t, 1)
    work = Path(tempfile.mkdtemp(prefix="render_diag_"))
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "render test"
    deck = work / "diag.pptx"
    prs.save(str(deck))
    t = time.time()
    try:
        pdf = to_pdf(deck, work, timeout)
        info["one_slide_pdf_seconds"] = round(time.time() - t, 1)
        info["pymupdf"] = load_pymupdf() is not None
        info["png_ok"] = bool(pdf_to_pngs(pdf, work, 640, [1]))
        info["result"] = "rendering works"
    except RenderError as e:
        info.update(stage=e.stage, error=e.error, result="rendering failed")
    return info


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck", nargs="?")
    ap.add_argument("--slides", help="which slides, e.g. 1-15 or 1,4,9-12 (default: all)")
    ap.add_argument("--out", help="output folder (default: <deck folder>/<deck name>_render)")
    ap.add_argument("--width", type=int, default=1280)
    ap.add_argument("--timeout", type=int, default=150, help="seconds allowed for LibreOffice (default 150)")
    ap.add_argument("--diagnose", action="store_true", help="test LibreOffice and image conversion on a 1-slide deck")
    ap.add_argument("--labels-from", help="for a PDF: the .pptx it came from, to label sheets with pattern IDs")
    a = ap.parse_args()

    if a.diagnose:
        print(json.dumps(diagnose(a.timeout), indent=2))
        return
    if not a.deck:
        ap.error("give a deck, or --diagnose")
    deck = Path(a.deck).resolve()
    out = Path(a.out) if a.out else deck.parent / f"{deck.stem}_render"
    out.mkdir(parents=True, exist_ok=True)
    started = time.time()
    try:
        if deck.suffix.lower() == ".pdf":
            fitz = load_pymupdf()
            if fitz is None:
                raise RenderError("images", "PyMuPDF could not be installed", "Open the PDF to check the slides.")
            with fitz.open(str(deck)) as doc:
                total = doc.page_count
            numbers = parse_slides(a.slides, total) if a.slides else list(range(1, total + 1))
            labels = pattern_labels(Path(a.labels_from) if a.labels_from else None, numbers)
            final_pdf = deck
            pngs = pdf_to_pngs(deck, out, a.width, numbers, pages=numbers)
        else:
            total = len(dk.open_deck(deck).slides)
            numbers = parse_slides(a.slides, total) if a.slides else list(range(1, total + 1))
            work = Path(tempfile.mkdtemp(prefix="render_"))
            src, labels = subset_deck(deck, numbers, work)
            pdf = to_pdf(src, work, a.timeout)
            final_pdf = out / f"{deck.stem}_{numbers[0]:03d}-{numbers[-1]:03d}.pdf"
            shutil.move(str(pdf), final_pdf)
            pngs = pdf_to_pngs(final_pdf, out, a.width, numbers)
        sheets = contact_sheets(pngs, labels, out, deck.stem)
    except RenderError as e:
        print(json.dumps({"ok": False, "stage": e.stage, "error": e.error, "hint": e.hint}, indent=2))
        sys.exit(1)
    remaining = None
    if a.slides and numbers[-1] < total:
        size = len(numbers)
        remaining = f"--slides {numbers[-1] + 1}-{min(total, numbers[-1] + size)}"
    print(json.dumps({"ok": True, "slides_rendered": f"{numbers[0]}-{numbers[-1]} of {total}",
                      "seconds": round(time.time() - started, 1), "pdf": str(final_pdf),
                      "contact_sheets": [str(s) for s in sheets], "slide_images": [str(p) for p in pngs],
                      "next_chunk": remaining}, indent=2))


if __name__ == "__main__":
    main()
