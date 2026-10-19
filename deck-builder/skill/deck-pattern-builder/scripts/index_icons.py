"""Index brand icons on a few PowerPoint slides so the builder can copy them into decks as vectors.

    python index_icons.py brand_icons.pptx --out DIR

Put the icon slides from the brand template in one file (icons and pictograms, ideally each with a
short text label under or beside it). For every icon (a group, freeform drawing or picture without
text) the script takes its name from the nearest label, or icon_s<slide>_<n> without one.

Writes to DIR:
  icons.pptx   a copy of the input with each icon shape named icon:<name> (Alt+F10 to check)
  icons.csv    name, keywords, slide, shape_id: the index the builder reads. Add synonyms to
               keywords (risk -> "risk compliance control audit") so plan words find the icon.
  icons.md     the icon names per slide, for the agent to choose from
Put icons.pptx and icons.csv in deck-builder/private/ and rebuild the skill zip.
"""
import argparse
import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import deckkit as dk  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from prep_patterns import abs_box  # noqa: E402

EMU_PER_IN = 914400
MIN_IN, MAX_IN = 0.12, 2.5
FIELDS = ["name", "keywords", "slide", "shape_id"]


def slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")
    return s[:40] or "icon"


class Box:
    def __init__(self, shp, parents, text=""):
        self.shp = shp
        self.x, self.y, self.w, self.h = abs_box(shp, parents)
        self.text = " ".join(text.split())
        self.label = None

    @property
    def cx(self):
        return self.x + self.w / 2

    @property
    def cy(self):
        return self.y + self.h / 2


def is_freeform(shp) -> bool:
    return shp._element.find(f".//{qn('a:custGeom')}") is not None


def icon_sized(b: Box) -> bool:
    if not (MIN_IN <= b.w <= MAX_IN and MIN_IN <= b.h <= MAX_IN):
        return False
    return 0.25 <= b.w / b.h <= 4


def collect(slide, slide_w, slide_h):
    icons, labels = [], []

    def visit(shapes, parents):
        for shp in shapes:
            text = dk.shape_text(shp).strip()
            is_group = shp.__class__.__name__ == "GroupShape"
            if is_group:
                b = Box(shp, parents)
                texts = [s for s, _ in dk.iter_shapes(shp.shapes) if dk.shape_text(s).strip()]
                if b.w > slide_w * 0.35 or b.h > slide_h * 0.35 or len(texts) > 1:
                    visit(shp.shapes, parents + (shp,))  # a container of several icons
                    continue
                if icon_sized(b):
                    if texts:  # icon grouped with its own label
                        b.label = " ".join(dk.shape_text(texts[0]).split())
                    icons.append(b)
                continue
            if text:
                if len(text) <= 60 and not (getattr(shp, "is_placeholder", False) and shp.is_placeholder):
                    labels.append(Box(shp, parents, text))
                continue
            if shp._element.tag == qn("p:pic") or is_freeform(shp):
                b = Box(shp, parents)
                if icon_sized(b):
                    icons.append(b)

    visit(slide.shapes, ())
    return icons, labels


def match_labels(icons, labels):
    pairs = []
    for i in icons:
        if i.label:
            continue
        for lab in labels:
            below = lab.y - (i.y + i.h)
            right = lab.x - (i.x + i.w)
            if -0.1 <= below < 0.9 and abs(lab.cx - i.cx) < max(i.w, 0.8):
                pairs.append((below, i, lab))
            elif -0.1 <= right < 1.2 and abs(lab.cy - i.cy) < max(i.h, 0.5) * 0.75:
                pairs.append((right + 0.05, i, lab))
    used = set()
    for _, i, lab in sorted(pairs, key=lambda p: p[0]):
        if i.label or id(lab) in used:
            continue
        i.label = lab.text
        used.add(id(lab))


def index(deck: Path):
    prs = dk.open_deck(deck)
    sw, sh = prs.slide_width / EMU_PER_IN, prs.slide_height / EMU_PER_IN
    rows, seen = [], {}
    for n, slide in enumerate(prs.slides, 1):
        icons, labels = collect(slide, sw, sh)
        match_labels(icons, labels)
        for k, i in enumerate(sorted(icons, key=lambda b: (round(b.y / 0.4), b.x)), 1):
            base = slug(i.label) if i.label else f"icon_s{n}_{k}"
            seen[base] = seen.get(base, 0) + 1
            name = base if seen[base] == 1 else f"{base}_{seen[base]}"
            i.shp.name = f"icon:{name}"
            rows.append({"name": name, "keywords": (i.label or "").lower(), "slide": n,
                         "shape_id": i.shp.shape_id})
    return prs, rows


def icons_md(rows) -> str:
    out = ["# Icons", "",
           "Icon names from the brand icon slides. Use a name, or a few keywords, in plan.json `icons`.", ""]
    by_slide = {}
    for r in rows:
        by_slide.setdefault(str(r["slide"]), []).append(r)
    for s, rs in by_slide.items():
        out.append(f"- Slide {s}: " + ", ".join(
            r["name"] + (f" ({r['keywords']})" if r.get("keywords") and r["keywords"] != r["name"].replace("_", " ")
                         else "") for r in rs))
    return "\n".join(out) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--out", default=None, help="output folder (default: next to the input)")
    a = ap.parse_args()
    src = Path(a.deck)
    out = Path(a.out) if a.out else src.parent
    out.mkdir(parents=True, exist_ok=True)
    prs, rows = index(src)
    prs.save(str(out / "icons.pptx"))
    with (out / "icons.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    (out / "icons.md").write_text(icons_md(rows), encoding="utf-8")
    unnamed = sum(1 for r in rows if r["name"].startswith("icon_s"))
    print(f"{len(rows)} icons indexed ({unnamed} without a label, named icon_s<slide>_<n>)")
    print(f"wrote {out / 'icons.pptx'}, {out / 'icons.csv'}, {out / 'icons.md'}")


if __name__ == "__main__":
    main()
