"""Shared helpers for the deck builder: slide copy/delete, tokens, pruning, icons, text fitting.

Works on python-pptx 1.0.x. Some helpers use python-pptx internals (`_sldIdLst`, `_element`)
because the public API has no slide duplication or deletion.
"""
from __future__ import annotations

import copy
import csv
import math
import os
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

TOKEN_RE = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")
ITEM_RE = re.compile(r"^(item\d+)(?:_|$)", re.I)
PATTERN_RE = re.compile(r"^\s*PATTERN\s*:\s*([A-Za-z0-9_\-]+)", re.I | re.M)
NOTE_FIELD_RE = re.compile(r"^[ \t]*([A-Z][A-Z _-]+?)[ \t]*:[ \t]*(.*)$", re.M)
EXAMPLE_RE = re.compile(r"^[ \t]*EXAMPLE[ \t]+([A-Za-z0-9_]+)[ \t]*:[ \t]*(.*)$", re.M)
TABLE_NOTE_RE = re.compile(r"^[ \t]*TABLE[ \t]+(t\d+)[ \t]*:[ \t]*(.*)$", re.M)
TABLE_TOKEN_RE = re.compile(r"^(t\d+)_r(\d+)c(\d+)$")
R_ATTRS = (qn("r:id"), qn("r:embed"), qn("r:link"), qn("r:pict"))
DEFAULT_FONT_PT = 18.0
EMU_PER_PT = 12700


# ---------------------------------------------------------------- paths

def skill_dir() -> Path:
    return Path(__file__).resolve().parent.parent


def default_out_dir() -> Path:
    d = Path("/mnt/data")
    return d if d.is_dir() and os.access(d, os.W_OK) else Path.cwd()


# ---------------------------------------------------------------- shapes

def iter_shapes(shapes, parents=()):
    """Yield (shape, parent_groups) for every shape, descending into groups."""
    for shp in shapes:
        yield shp, parents
        if shp.shape_type == 6 or shp.__class__.__name__ == "GroupShape":  # MSO_SHAPE_TYPE.GROUP
            yield from iter_shapes(shp.shapes, parents + (shp,))


def text_frames(shape):
    """Text frames in a shape, including table cells."""
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        yield shape.text_frame
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


def shape_text(shape) -> str:
    return "\n".join(tf.text for tf in text_frames(shape))


def tokens_in(shape) -> list[str]:
    return TOKEN_RE.findall(shape_text(shape))


def slide_tokens(slide) -> list[str]:
    seen = []
    for shp, _ in iter_shapes(slide.shapes):
        for t in tokens_in(shp):
            if t not in seen:
                seen.append(t)
    return seen


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


# ---------------------------------------------------------------- notes / pattern metadata

def notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text or ""


def pattern_id(slide) -> str | None:
    m = PATTERN_RE.search(notes_text(slide))
    return m.group(1).upper() if m else None


def pattern_meta(slide) -> dict:
    """All 'KEY: value' lines in the notes, e.g. PATTERN, PURPOSE, USE WHEN, ITEMS."""
    return {k.strip().upper(): v.strip() for k, v in NOTE_FIELD_RE.findall(notes_text(slide)) if v.strip()}


def pattern_examples(slide) -> dict:
    """'EXAMPLE slot: text' lines from the notes: the human text each slot replaced."""
    return {k: v.strip() for k, v in EXAMPLE_RE.findall(notes_text(slide))}


def table_dims(slide) -> dict:
    """Tables with tokenised cells: {'t1': (rows, cols)} from the cell tokens on the slide."""
    dims = {}
    for t in slide_tokens(slide):
        m = TABLE_TOKEN_RE.match(t)
        if m:
            r, c = int(m.group(2)), int(m.group(3))
            pr, pc = dims.get(m.group(1), (0, 0))
            dims[m.group(1)] = (max(pr, r), max(pc, c))
    return dims


def set_notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text or ""


# ---------------------------------------------------------------- slide operations

def duplicate_slide(prs, src):
    """Copy `src` to a new slide at the end of the deck (same layout, shapes, images, background)."""
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        remove_shape(shp)

    rid_map = {}
    for rel in src.part.rels.values():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rel.rId] = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rid_map[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)

    src_tree, new_tree = src.shapes._spTree, new.shapes._spTree
    for el in src_tree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        new_tree.append(copy.deepcopy(el))

    src_bg = src._element.cSld.find(qn("p:bg"))
    if src_bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(src_bg))

    for el in new._element.iter():
        for attr in R_ATTRS:
            old = el.get(attr)
            if old in rid_map:
                el.set(attr, rid_map[old])
    return new


def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
            return
    raise ValueError("slide not in presentation")


def move_slide(prs, slide, new_index: int):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            sld_ids.remove(sld_id)
            sld_ids.insert(new_index, sld_id)
            return
    raise ValueError("slide not in presentation")


def index_patterns(prs) -> dict:
    out = {}
    for i, s in enumerate(prs.slides):
        pid = pattern_id(s)
        if pid:
            if pid in out:
                raise ValueError(f"pattern {pid} appears twice (slides {out[pid][0] + 1} and {i + 1})")
            out[pid] = (i, s)
    return out


# ---------------------------------------------------------------- token filling

def _as_lines(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [str(v) for v in value if str(v).strip()]
    s = str(value)
    return [line for line in s.split("\n")] if s.strip() else []


def _merge_runs(paragraph):
    """Put all run text into the first run (keeps its formatting), drop the other runs."""
    runs = paragraph.runs
    if len(runs) < 2:
        return
    runs[0].text = "".join(r.text for r in runs)
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _fill_paragraph(paragraph, values: dict, filled: set, unfilled: set):
    """Replace tokens in one paragraph. Returns extra lines when the paragraph is a lone list token."""
    full = "".join(r.text for r in paragraph.runs)
    toks = TOKEN_RE.findall(full)
    if not toks:
        return None
    # A paragraph holding only one token whose value has several lines becomes several paragraphs.
    only = TOKEN_RE.fullmatch(full.strip())
    if only:
        lines = _as_lines(values.get(only.group(1)))
        if len(lines) > 1:
            filled.add(only.group(1))
            _merge_runs(paragraph)
            paragraph.runs[0].text = lines[0]
            return lines[1:]
    # Tokens inside a single run keep that run's formatting; tokens split across runs force a merge.
    if any(TOKEN_RE.search(r.text) is None for r in paragraph.runs if "{{" in r.text or "}}" in r.text):
        _merge_runs(paragraph)
    for run in paragraph.runs:
        def sub(m):
            name = m.group(1)
            lines = _as_lines(values.get(name))
            if lines:
                filled.add(name)
                return " ".join(lines)
            unfilled.add(name)
            return m.group(0)
        run.text = TOKEN_RE.sub(sub, run.text)
    return None


def expand_tables(values: dict) -> tuple[dict, dict]:
    """Turn table values (`"t1": [[row1 cells], [row2 cells]]`) into per-cell tokens (`t1_r1c1`).

    Returns (values, rows_given) where rows_given = {"t1": number of rows supplied}. Rows the plan
    does not supply are removed from the table; cells supplied as "" are blanked.
    """
    out, rows_given = dict(values), {}
    for k, v in values.items():
        if re.fullmatch(r"t\d+", k) and isinstance(v, list):
            rows_given[k] = len(v)
            for r, row in enumerate(v, 1):
                for c, cell in enumerate(row if isinstance(row, list) else [row], 1):
                    if cell is not None and str(cell).strip():
                        out[f"{k}_r{r}c{c}"] = cell
            del out[k]
    return out, rows_given


def fill_tokens(slide, values: dict) -> tuple[set, set]:
    """Fill {{tokens}} on a slide. Returns (filled, unfilled) token names. Unfilled tokens stay as text."""
    values, _ = expand_tables(values)
    filled, unfilled = set(), set()
    for shp, _ in iter_shapes(slide.shapes):
        for tf in text_frames(shp):
            for p in list(tf.paragraphs):
                extra = _fill_paragraph(p, values, filled, unfilled)
                anchor = p._p
                for line in extra or []:
                    new_p = copy.deepcopy(anchor)
                    for r in new_p.findall(qn("a:r"))[1:]:
                        new_p.remove(r)
                    new_p.find(qn("a:r")).find(qn("a:t")).text = line
                    anchor.addnext(new_p)
                    anchor = new_p
    return filled, unfilled - filled


def _prune_table_rows(shp, unfilled: set, rows_given: dict) -> int:
    """Remove table rows whose tokens were all left unfilled and that the plan did not supply."""
    tbl = shp.table._tbl
    if any(tc.get("rowSpan") not in (None, "1") or tc.get("vMerge") for tc in tbl.iter(qn("a:tc"))):
        return 0  # merged rows: removing one would break the table, so only blank
    gone = 0
    for r_idx, tr in enumerate(list(tbl.tr_lst), 1):
        toks = TOKEN_RE.findall("".join(t.text or "" for t in tr.iter(qn("a:t"))))
        if not toks or not all(t in unfilled for t in toks):
            continue
        m = TABLE_TOKEN_RE.match(toks[0])
        given = rows_given.get(m.group(1)) if m else None
        if given is None or r_idx > given:
            tr.getparent().remove(tr)
            gone += 1
    if gone:
        shp.height = sum(row.height for row in shp.table.rows)
    return gone


def prune_unused(slide, filled: set, unfilled: set, rows_given: dict | None = None) -> list[str]:
    """Remove empty slots. Returns a list of what was removed or blanked.

    - Table rows with only unfilled tokens, beyond the rows the plan supplied: removed.
    - An item (token prefix `itemN_`) with no filled tokens: remove every shape or group named `itemN...`
      and every shape whose tokens all belong to that item.
    - Any other shape whose text is only unfilled tokens: removed.
    - Unfilled tokens left inside other text: blanked.
    """
    removed = []
    items_all = {m.group(1).lower() for t in filled | unfilled if (m := ITEM_RE.match(t))}
    items_used = {m.group(1).lower() for t in filled if (m := ITEM_RE.match(t))}
    dead_items = items_all - items_used

    def item_of(name):
        m = ITEM_RE.match(name or "")
        return m.group(1).lower() if m else None

    for shp, _ in list(iter_shapes(slide.shapes)):
        if not (getattr(shp, "has_table", False) and shp.has_table):
            continue
        names = {m.group(1) for t in tokens_in(shp) if (m := TABLE_TOKEN_RE.match(t))}
        if names & set(rows_given or {}):
            gone = _prune_table_rows(shp, unfilled, rows_given or {})
            if gone:
                removed.append(f"removed {gone} unused row(s) from table {shp.name}")

    for shp, parents in list(iter_shapes(slide.shapes)):
        if any(p._element.getparent() is None for p in parents) or shp._element.getparent() is None:
            continue  # already removed with its group
        it = item_of(shp.name)
        toks = tokens_in(shp)
        tok_items = {item_of(t) for t in toks}
        if it in dead_items or (toks and tok_items <= dead_items and None not in tok_items):
            removed.append(f"removed {shp.name}")
            remove_shape(shp)
            continue
        if toks and all(t in unfilled for t in toks) and not TOKEN_RE.sub("", shape_text(shp)).strip():
            removed.append(f"removed {shp.name} (empty: {', '.join(toks)})")
            remove_shape(shp)
            continue
        if any(t in unfilled for t in toks):
            for tf in text_frames(shp):
                for p in tf.paragraphs:
                    for r in p.runs:
                        r.text = TOKEN_RE.sub(lambda m: "" if m.group(1) in unfilled else m.group(0), r.text)
            removed.append(f"blanked {', '.join(t for t in toks if t in unfilled)} in {shp.name}")
    return removed


# ---------------------------------------------------------------- icons

def load_icon_index(icons_dir: Path) -> list[dict]:
    idx = icons_dir / "icons_index.csv"
    if not idx.is_file():
        return [{"file": p.name, "name": p.stem, "keywords": p.stem.replace("_", " ")}
                for p in sorted(icons_dir.glob("*.png"))]
    with idx.open(newline="", encoding="utf-8-sig") as f:
        return [{k.strip().lower(): (v or "").strip() for k, v in row.items()} for row in csv.DictReader(f)]


def find_icon(query: str, icons_dir: Path, index: list[dict] | None = None) -> Path | None:
    """Exact file or name match first, then the best keyword overlap."""
    index = index if index is not None else load_icon_index(icons_dir)
    q = query.strip().lower()
    words = set(re.findall(r"[a-z0-9]+", q))
    best, best_score = None, 0
    for row in index:
        name = row.get("name", "").lower()
        if q in (row.get("file", "").lower(), name):
            return icons_dir / row["file"]
        kw = set(re.findall(r"[a-z0-9]+", (row.get("keywords", "") + " " + name).lower()))
        score = len(words & kw)
        if score > best_score:
            best, best_score = row, score
    return icons_dir / best["file"] if best else None


def find_shape(slide, name: str):
    for shp, _ in iter_shapes(slide.shapes):
        if shp.name == name:
            return shp
    return None


def set_icon(slide, shape_name: str, png: Path | None, query: str) -> str:
    shp = find_shape(slide, shape_name)
    if shp is None:
        return f"no shape named {shape_name}"
    if png is None or not png.is_file():
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            shp.text_frame.text = f"[ICON: {query}]"
            return f"[ICON: {query}] placeholder in {shape_name}"
        return f"icon '{query}' not found; {shape_name} left as is"
    blip = shp._element.find(".//" + qn("a:blip"))
    if blip is not None and shp._element.tag == qn("p:pic"):  # picture: swap the image, keep the frame
        _, rid = slide.part.get_or_add_image_part(str(png))
        blip.set(qn("r:embed"), rid)
        return f"{shape_name} ← {png.name}"
    # Not a picture: put the picture in the same box, inside the same group, then drop the stand-in.
    parent_el = shp._element.getparent()
    pic = slide.shapes.add_picture(str(png), shp.left, shp.top, shp.width, shp.height)
    pic.name = shape_name
    parent_el.insert(parent_el.index(shp._element), pic._element)
    remove_shape(shp)
    return f"{shape_name} ← {png.name} (replaced stand-in shape)"


# ---------------------------------------------------------------- text sizing

def _sz_from(el) -> float | None:
    if el is None:
        return None
    sz = el.get("sz")
    return int(sz) / 100 if sz else None


def effective_font_pt(shape, run=None, paragraph=None) -> float:
    if run is not None and run.font.size:
        return run.font.size.pt
    if paragraph is not None:
        ppr = paragraph._p.find(qn("a:pPr"))
        if ppr is not None and (s := _sz_from(ppr.find(qn("a:defRPr")))):
            return s
        end = paragraph._p.find(qn("a:endParaRPr"))
        if (s := _sz_from(end)):
            return s
    tx = shape._element.find(".//" + qn("p:txBody"))
    if tx is not None:
        lvl = tx.find(f"{qn('a:lstStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
        if (s := _sz_from(lvl)):
            return s
    if getattr(shape, "is_placeholder", False) and shape.is_placeholder:
        return _placeholder_font_pt(shape)
    return DEFAULT_FONT_PT


def _placeholder_font_pt(shape) -> float:
    """Size inherited by a slide placeholder: layout placeholder first, then the master's text styles."""
    is_title = shape.placeholder_format.type in (1, 3)  # TITLE, CENTER_TITLE
    try:
        layout = shape.part.slide_layout
    except AttributeError:
        return 36.0 if is_title else DEFAULT_FONT_PT
    try:
        lph = layout.placeholders.get(idx=shape.placeholder_format.idx)
    except Exception:  # noqa: BLE001
        lph = None
    if lph is not None:
        tx = lph._element.find(".//" + qn("p:txBody"))
        if tx is not None:
            for el in tx.iter(qn("a:defRPr"), qn("a:rPr"), qn("a:endParaRPr")):
                if (s := _sz_from(el)):
                    return s
    styles = layout.slide_master._element.find(qn("p:txStyles"))
    if styles is not None:
        node = styles.find(qn("p:titleStyle") if is_title else qn("p:bodyStyle"))
        if node is not None:
            el = node.find(f"{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
            if (s := _sz_from(el)):
                return s
    return 36.0 if is_title else DEFAULT_FONT_PT


def estimate_lines(text: str, width_pt: float, font_pt: float) -> int:
    chars_per_line = max(1, int(width_pt / (font_pt * 0.5)))
    return sum(max(1, math.ceil(len(line) / chars_per_line)) for line in text.split("\n"))


def text_fit(shape) -> dict | None:
    """Rough check: does the text fit the box? Returns needed vs available height in points."""
    if not (getattr(shape, "has_text_frame", False) and shape.has_text_frame) or not shape.text_frame.text.strip():
        return None
    if shape.width is None or shape.height is None:
        return None
    tf = shape.text_frame
    margin_w = ((tf.margin_left or Emu(91440)) + (tf.margin_right or Emu(91440))) / EMU_PER_PT
    margin_h = ((tf.margin_top or Emu(45720)) + (tf.margin_bottom or Emu(45720))) / EMU_PER_PT
    width = shape.width / EMU_PER_PT - margin_w
    needed = 0.0
    sizes = []
    for p in tf.paragraphs:
        size = effective_font_pt(shape, p.runs[0] if p.runs else None, p)
        sizes.append(size)
        txt = "".join(r.text for r in p.runs)
        needed += estimate_lines(txt, width, size) * size * 1.2
    return {"needed_pt": round(needed + margin_h, 1), "available_pt": round(shape.height / EMU_PER_PT, 1),
            "font_pt": max(sizes) if sizes else DEFAULT_FONT_PT}


def fit_text(shape, max_shrink_pt: float = 2.0) -> str | None:
    """Shrink text by up to `max_shrink_pt` if the estimate says it overflows. Returns a note or None."""
    fit = text_fit(shape)
    if not fit or fit["needed_pt"] <= fit["available_pt"]:
        return None
    tf = shape.text_frame
    runs = [(r, r.font.size, effective_font_pt(shape, r, p)) for p in tf.paragraphs for r in p.runs]
    for step in (1, 2):
        if step > max_shrink_pt:
            break
        for r, _, size in runs:
            r.font.size = Pt(size - step)
        fit = text_fit(shape)
        if fit["needed_pt"] <= fit["available_pt"]:
            return f"{shape.name}: shrank text by {step}pt to fit"
    for r, explicit, _ in runs:  # still too long: restore the original sizes and report it
        r.font.size = explicit
    return (f"{shape.name}: OVERFLOW, needs ~{fit['needed_pt']:.0f}pt of height, box has {fit['available_pt']:.0f}pt."
            " Shorten the text.")


def capacity_words(shape) -> int | None:
    """Rough number of words a text box holds at its current font size."""
    if shape.width is None or shape.height is None:
        return None
    p = shape.text_frame.paragraphs[0] if shape.text_frame.paragraphs else None
    size = effective_font_pt(shape, p.runs[0] if p is not None and p.runs else None, p)
    width = shape.width / EMU_PER_PT - 14
    height = shape.height / EMU_PER_PT - 7
    lines = max(1, int(height / (size * 1.2)))
    chars = lines * max(1, int(width / (size * 0.5)))
    return max(1, chars // 6)


# ---------------------------------------------------------------- theme

def theme_fonts(prs) -> set[str]:
    fonts = set()
    try:
        theme = prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
        for tag in ("a:majorFont", "a:minorFont"):
            for latin in root.iter(qn(tag)):
                el = latin.find(qn("a:latin"))
                if el is not None and el.get("typeface"):
                    fonts.add(el.get("typeface"))
    except Exception:  # noqa: BLE001
        pass
    return fonts


def open_deck(path) -> Presentation:
    return Presentation(str(path))
