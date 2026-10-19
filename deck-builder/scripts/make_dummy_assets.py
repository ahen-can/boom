"""Create brand-neutral dummy assets for the skill: assets/patterns.pptx and assets/icons/.

These stand in for the real PwC patterns and icons, which stay on the user's laptop.
Run: python deck-builder/scripts/make_dummy_assets.py
"""
import csv
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SKILL = Path(__file__).resolve().parent.parent / "skill" / "deck-pattern-builder"
ASSETS = SKILL / "assets"
ICONS = ASSETS / "icons"

INK = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x7F, 0x7F, 0x7F)
ACCENT = RGBColor(0x2F, 0x5D, 0x7C)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_HEX = (0x2F, 0x5D, 0x7C, 255)


# ---------------------------------------------------------------- icons

def draw_icon(name):
    img = Image.new("RGBA", (256, 256), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    c, w = ACCENT_HEX, 18
    if name == "growth":
        d.line([(30, 200), (100, 130), (150, 170), (226, 70)], fill=c, width=w)
        d.polygon([(226, 40), (236, 100), (186, 70)], fill=c)
    elif name == "risk":
        d.polygon([(128, 25), (235, 225), (21, 225)], outline=c, width=w)
        d.line([(128, 95), (128, 160)], fill=c, width=w)
        d.ellipse([(118, 180), (138, 200)], fill=c)
    elif name == "people":
        for x in (80, 176):
            d.ellipse([(x - 35, 40), (x + 35, 110)], outline=c, width=w)
            d.arc([(x - 60, 120), (x + 60, 260)], 180, 360, fill=c, width=w)
    elif name == "data":
        for i, h in enumerate((80, 140, 190)):
            d.rectangle([(40 + i * 65, 230 - h), (85 + i * 65, 230)], fill=c)
    elif name == "check":
        d.ellipse([(20, 20), (236, 236)], outline=c, width=w)
        d.line([(70, 130), (115, 175), (190, 85)], fill=c, width=w)
    elif name == "target":
        for r in (108, 70, 32):
            d.ellipse([(128 - r, 128 - r), (128 + r, 128 + r)], outline=c, width=14)
    img.save(ICONS / f"{name}.png")


ICON_KEYWORDS = {
    "growth": "growth revenue increase trend up performance",
    "risk": "risk warning issue compliance alert",
    "people": "people team workforce talent stakeholders customers",
    "data": "data analytics reporting metrics insight",
    "check": "check done quality assurance complete approved",
    "target": "target objective goal focus outcome",
}


def make_icons():
    ICONS.mkdir(parents=True, exist_ok=True)
    for name in ICON_KEYWORDS:
        draw_icon(name)
    with (ICONS / "icons_index.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["file", "name", "keywords"])
        for name, kw in ICON_KEYWORDS.items():
            w.writerow([f"{name}.png", name, kw])


# ---------------------------------------------------------------- pattern helpers

def text(shapes, name, x, y, w, h, value, size=12, bold=False, color=INK, align=None, anchor=MSO_ANCHOR.TOP):
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return tb


def box(shapes, name, kind, x, y, w, h, fill, line=None):
    s = shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def header(slide, tracker=True):
    if tracker:
        text(slide.shapes, "tracker", 0.5, 0.2, 8, 0.3, "{{tracker}}", size=10, color=GREY)
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.5), Inches(0.5), Inches(12.3), Inches(0.9)
    t.text_frame.text = "{{title}}"
    t.text_frame.word_wrap = True
    t.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    r = t.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb = Pt(24), True, INK
    t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT


def footer(slide):
    text(slide.shapes, "source", 0.5, 6.9, 10, 0.3, "{{source}}", size=9, color=GREY)


def notes(slide, pid, purpose, use_when, items=None):
    lines = [f"PATTERN: {pid}", f"PURPOSE: {purpose}", f"USE WHEN: {use_when}"]
    if items:
        lines.append(f"ITEMS: {items}")
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


# ---------------------------------------------------------------- patterns

def p_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s.shapes, "band", MSO_SHAPE.RECTANGLE, 0, 0, 0.35, 7.5, ACCENT)
    text(s.shapes, "title", 1.0, 2.4, 11, 1.5, "{{title}}", size=40, bold=True)
    text(s.shapes, "subtitle", 1.0, 4.0, 11, 0.8, "{{subtitle}}", size=20, color=GREY)
    text(s.shapes, "date", 1.0, 6.4, 6, 0.4, "{{date}}", size=12, color=GREY)
    notes(s, "TITLE", "Cover slide", "First slide of every deck")


def p_section(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s.shapes, "panel", MSO_SHAPE.RECTANGLE, 0, 0, 4.5, 7.5, ACCENT)
    text(s.shapes, "section_number", 0.8, 2.6, 3, 1.5, "{{section_number}}", size=72, bold=True, color=WHITE)
    text(s.shapes, "section_title", 5.2, 3.0, 7.5, 1.5, "{{section_title}}", size=32, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "SECTION", "Section divider", "Start of each major section (3-6 per deck)")


def p_obj_3col(prs, icon_png):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    header(s)
    for i in range(3):
        n, x = i + 1, 0.5 + i * 4.15
        g = s.shapes.add_group_shape()
        g.name = f"item{n}"
        box(g.shapes, f"item{n}_bg", MSO_SHAPE.RECTANGLE, x, 1.7, 3.9, 4.9, LIGHT)
        pic = g.shapes.add_picture(str(icon_png), Inches(x + 0.3), Inches(1.95), Inches(0.7), Inches(0.7))
        pic.name = f"item{n}_icon"
        text(g.shapes, f"item{n}_head", x + 0.3, 2.85, 3.3, 0.8, f"{{{{item{n}_head}}}}", size=16, bold=True,
             color=ACCENT)
        text(g.shapes, f"item{n}_body", x + 0.3, 3.7, 3.3, 2.7, f"{{{{item{n}_body}}}}", size=12)
    footer(s)
    notes(s, "OBJ_3COL", "Three objectives or themes side by side, each with an icon",
          "Objectives, key themes, pillars when there are exactly 2-3 parallel ideas", "2-3")


def p_obj_numbered_4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    header(s)
    for i in range(4):
        n, y = i + 1, 1.7 + i * 1.25
        g = s.shapes.add_group_shape()
        g.name = f"item{n}"
        c = box(g.shapes, f"item{n}_num", MSO_SHAPE.OVAL, 0.5, y + 0.15, 0.6, 0.6, ACCENT)
        c.text_frame.text = str(n)
        r = c.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb = Pt(16), True, WHITE
        c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        text(g.shapes, f"item{n}_head", 1.4, y + 0.1, 3.6, 0.9, f"{{{{item{n}_head}}}}", size=14, bold=True)
        text(g.shapes, f"item{n}_body", 5.2, y + 0.1, 7.6, 0.9, f"{{{{item{n}_body}}}}", size=12)
        box(g.shapes, f"item{n}_rule", MSO_SHAPE.RECTANGLE, 1.4, y + 1.1, 11.4, 0.02, LIGHT)
    footer(s)
    notes(s, "OBJ_NUMBERED_4", "Numbered list of objectives or priorities with a short explanation each",
          "Objectives or priorities when there are 3-4 items needing a sentence each", "3-4")


def p_chevron_4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    header(s)
    for i in range(4):
        n, x = i + 1, 0.5 + i * 3.1
        g = s.shapes.add_group_shape()
        g.name = f"item{n}"
        ch = box(g.shapes, f"item{n}_chevron", MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, x, 1.7, 3.2, 0.9,
                 ACCENT)
        ch.text_frame.text = f"{{{{item{n}_phase}}}}"
        r = ch.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb = Pt(14), True, WHITE
        ch.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        text(g.shapes, f"item{n}_time", x + 0.2, 2.7, 2.7, 0.35, f"{{{{item{n}_time}}}}", size=10, color=GREY)
        text(g.shapes, f"item{n}_body", x + 0.2, 3.1, 2.7, 3.5, f"{{{{item{n}_body}}}}", size=12)
    footer(s)
    notes(s, "APPROACH_CHEVRON_4", "Phased approach, left to right, with activities under each phase",
          "Approach, methodology or roadmap with 3-4 sequential phases", "3-4")


def p_two_col(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    header(s)
    text(s.shapes, "left_head", 0.5, 1.7, 5.9, 0.6, "{{left_head}}", size=16, bold=True, color=ACCENT)
    text(s.shapes, "left_body", 0.5, 2.4, 5.9, 4.2, "{{left_body}}", size=13)
    box(s.shapes, "divider", MSO_SHAPE.RECTANGLE, 6.65, 1.7, 0.02, 4.9, GREY)
    text(s.shapes, "right_head", 6.9, 1.7, 5.9, 0.6, "{{right_head}}", size=16, bold=True, color=ACCENT)
    text(s.shapes, "right_body", 6.9, 2.4, 5.9, 4.2, "{{right_body}}", size=13)
    footer(s)
    notes(s, "TWO_COL", "Two contrasting or complementary blocks",
          "Current vs future state, challenge vs response, what we heard vs what it means")


def p_table_matrix(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    header(s)
    rows, cols = 7, 4
    frame = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.7), Inches(12.3), Inches(0.5 * rows))
    frame.name = "t1"
    tbl = frame.table
    tbl.columns[0].width = Inches(5.4)
    for c in range(1, cols):
        tbl.columns[c].width = Inches(2.3)
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = f"{{{{t1_r{r + 1}c{c + 1}}}}}"
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(12 if r == 0 else 11)
            run.font.bold = r == 0
    footer(s)
    notes(s, "GOV_MATRIX", "Matrix of activities against forums, owners or phases, with ticks",
          "Governance, RACI or activity-by-phase views with up to 6 rows and 3 columns")
    s.notes_slide.notes_text_frame.text += "\nTABLE t1: 7 rows x 4 columns"


def make_patterns():
    prs = Presentation()
    k = 13.333 / 10
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for container in [prs.slide_master, *prs.slide_layouts]:
        for ph in container.placeholders:
            ph.left, ph.width = int(ph.left * k), int(ph.width * k)
    icon = ICONS / "target.png"
    p_title(prs)
    p_section(prs)
    p_obj_3col(prs, icon)
    p_obj_numbered_4(prs)
    p_chevron_4(prs)
    p_two_col(prs)
    p_table_matrix(prs)
    prs.save(str(ASSETS / "patterns.pptx"))


if __name__ == "__main__":
    make_icons()
    make_patterns()
    print(f"wrote {ASSETS / 'patterns.pptx'} and {len(ICON_KEYWORDS)} icons")
