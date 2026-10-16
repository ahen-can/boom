"""Build the pptx-probe skill: test deck, marker, SKILL.md (with probe.py inlined) and the upload zips.

Run from anywhere:  python deck-builder/scripts/make_probe_zip.py
Needs python-pptx (only to create assets/probe.pptx).
"""
import zipfile
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "probe-skill" / "pptx-probe"
DIST = ROOT / "dist"

SKILL_MD = """---
name: pptx-probe
description: Diagnostic skill. Use only when the user asks to run the pptx probe. Runs a Python script that reports what the code environment can do for PowerPoint generation (bundled skill files, python-pptx, LibreOffice, network, file output).
---

# pptx probe

Diagnostic only. Do not build a presentation or give advice. Run the probe and report.

## Steps

1. Locate this skill's bundled files in the code environment. Look for `scripts/probe.py` and
   `assets/marker.txt` (search the filesystem if needed, e.g. `find / -name marker.txt -path "*assets*" 2>/dev/null`).
2. If you found `scripts/probe.py`, run it with Python exactly as it is.
   If you could not find it, run the inline copy below instead, unchanged, and say clearly that the
   bundled files were NOT found.
3. Show the user the complete JSON output in a code block. Do not summarise or trim it.
4. Offer the files it wrote (`probe_xml.pptx`, `probe_pptx.pptx`) for download. They are in the
   folder given by `output_folder` in the report.
5. Below the JSON, add one line per item in `summary`: the key and yes/no. Nothing else.

## Inline copy of scripts/probe.py

```python
{probe}
```
"""


def make_assets():
    assets = SKILL / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    (assets / "marker.txt").write_text("pptx-probe-marker-7f3a\n")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "{{probe_token}}"
    slide.placeholders[1].text = "If the title above changed, direct XML editing works in this environment."
    prs.save(str(assets / "probe.pptx"))


def make_skill_md():
    probe = (SKILL / "scripts" / "probe.py").read_text().rstrip()
    (SKILL / "SKILL.md").write_text(SKILL_MD.replace("{probe}", probe))


def make_zips():
    DIST.mkdir(exist_ok=True)
    files = sorted(p for p in SKILL.rglob("*") if p.is_file() and "__pycache__" not in p.parts)
    # Main zip: files inside a pptx-probe/ folder (the usual skill bundle layout).
    with zipfile.ZipFile(DIST / "pptx-probe.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for p in files:
            z.write(p, Path("pptx-probe") / p.relative_to(SKILL))
    # Fallback zip: SKILL.md at the zip root, in case the uploader expects that.
    with zipfile.ZipFile(DIST / "pptx-probe-flat.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for p in files:
            z.write(p, p.relative_to(SKILL))


if __name__ == "__main__":
    make_assets()
    make_skill_md()
    make_zips()
    for z in sorted(DIST.glob("*.zip")):
        print(f"{z.relative_to(ROOT.parent)}  {z.stat().st_size / 1024:.1f} KB")
